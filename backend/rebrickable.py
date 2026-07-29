"""Rebrickable API client + parsing des grandes listes (Excel / CSV)."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

API_BASE = "https://rebrickable.com/api/v3"
SET_NUM_RE = re.compile(r"^\d+(?:-\d+)?$", re.IGNORECASE)


def normalize_set_num(raw: str, version: str | None = None) -> str | None:
    value = (raw or "").strip().strip('"').upper()
    if not value:
        return None

    # Strip accidental wrapping quotes / spaces
    value = value.replace(" ", "")

    if SET_NUM_RE.match(value):
        if "-" in value:
            return value
        # Attach version if provided (e.g. 75192 + 1 → 75192-1)
        if version and str(version).strip().isdigit():
            return f"{value}-{int(str(version).strip())}"
        return f"{value}-1"

    match = re.search(r"(\d+(?:-\d+)?)", value)
    if not match:
        return None
    found = match.group(1).upper()
    if "-" in found:
        return found
    if version and str(version).strip().isdigit():
        return f"{found}-{int(str(version).strip())}"
    return f"{found}-1"


def _parse_qty(raw: Any) -> int:
    try:
        return max(1, int(float(str(raw).strip())))
    except (TypeError, ValueError):
        return 1


def _expand_cell(value: Any) -> list[str]:
    """Une cellule peut contenir '75192-1,1,2' en entier."""
    if value is None:
        return []
    text = str(value).strip()
    if not text or text.lower() in {"none", "nan"}:
        return []
    # Excel may store numbers as floats (75192.0)
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int):
        text = str(value)
    if "," in text or ";" in text:
        return [p.strip() for p in re.split(r"[,;]", text) if p.strip()]
    return [text]


def _parse_fields(fields: list[Any]) -> dict[str, Any] | None:
    """
    Formats acceptés par ligne / cellule :
    - 75372-1,13,1  → set 75372-1, quantité 13, version 1 (ignorée)
    - 75372,13,1    → set 75372-1, quantité 13
    - 75372-1,2     → set 75372-1, quantité 2
    - 75372-1       → set 75372-1, quantité 1
    """
    expanded: list[str] = []
    for field in fields:
        expanded.extend(_expand_cell(field))
    cleaned = [f.strip().strip('"') for f in expanded if str(f).strip() != ""]
    if not cleaned:
        return None

    if len(cleaned) >= 3:
        # set, quantity, version
        raw_set, raw_qty, raw_ver = cleaned[0], cleaned[1], cleaned[2]
        if "-" in raw_set:
            set_num = normalize_set_num(raw_set)
        else:
            set_num = normalize_set_num(raw_set, raw_ver)
        qty = _parse_qty(raw_qty)
    elif len(cleaned) == 2:
        set_num = normalize_set_num(cleaned[0])
        qty = _parse_qty(cleaned[1])
    else:
        set_num = normalize_set_num(cleaned[0])
        qty = 1

    if not set_num:
        return None
    return {"set_num": set_num, "quantity": qty}


def _collect_from_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    by_num: dict[str, dict[str, Any]] = {}

    def add(item: dict[str, Any] | None) -> None:
        if not item:
            return
        key = item["set_num"]
        if key in by_num:
            by_num[key]["quantity"] = int(by_num[key]["quantity"]) + int(
                item["quantity"]
            )
        else:
            by_num[key] = dict(item)

    if not rows:
        return []

    first = [str(c).strip() if c is not None else "" for c in rows[0]]
    header_l = [h.lower() for h in first]
    looks_like_header = any(
        h in {
            "set number",
            "set_num",
            "setnumber",
            "set",
            "quantity",
            "qty",
            "number",
            "set id",
            "version",
        }
        for h in header_l
    )

    data_rows = rows[1:] if looks_like_header else rows
    set_idx = 0
    version_idx = None
    qty_idx = None
    if looks_like_header:
        for i, h in enumerate(header_l):
            if h in {
                "set number",
                "set_num",
                "setnumber",
                "set",
                "number",
                "set id",
                "set_id",
            }:
                set_idx = i
            elif h in {"version", "variant", "ver"}:
                version_idx = i
            elif h in {"quantity", "qty", "owned", "count"}:
                qty_idx = i

    for row in data_rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        if looks_like_header and qty_idx is not None:
            raw_set = row[set_idx] if set_idx < len(row) else None
            version = (
                row[version_idx]
                if version_idx is not None and version_idx < len(row)
                else None
            )
            qty_raw = row[qty_idx] if qty_idx < len(row) else 1
            set_num = normalize_set_num(str(raw_set or ""), version)
            if set_num:
                add({"set_num": set_num, "quantity": _parse_qty(qty_raw)})
            continue
        add(_parse_fields(list(row)))

    return list(by_num.values())


def parse_sets_csv(content: str | bytes) -> list[dict[str, Any]]:
    if isinstance(content, bytes):
        text = content.decode("utf-8-sig")
    else:
        text = content.lstrip("\ufeff")

    text = text.strip()
    if not text:
        raise ValueError("Fichier vide.")

    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    stream = io.StringIO(text)
    reader = csv.reader(stream, dialect)
    rows = [list(row) for row in reader]
    sets = _collect_from_rows(rows)
    if not sets:
        raise ValueError(
            "Aucun numéro de set trouvé. "
            "Attendu : 75372-1,13,1 (set, quantité, version) ou une colonne Set Number"
        )
    return sets


def parse_sets_excel(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    rows: list[list[Any]] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
    wb.close()
    sets = _collect_from_rows(rows)
    if not sets:
        raise ValueError(
            "Aucun numéro de set trouvé dans l’Excel. "
            "Format : 75372-1,13,1 (set, quantité, version)."
        )
    return sets


def parse_sets_file(content: bytes, filename: str | None = None) -> list[dict[str, Any]]:
    """Parse une grande liste thématique (Film, Série, Jeux…) depuis Excel ou CSV."""
    name = (filename or "").lower()
    suffix = Path(name).suffix

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return parse_sets_excel(content)
    if suffix in {".csv", ".txt"}:
        return parse_sets_csv(content)

    # Sniff by magic bytes (zip/xlsx starts with PK)
    if content[:2] == b"PK":
        return parse_sets_excel(content)
    try:
        return parse_sets_csv(content)
    except (UnicodeDecodeError, ValueError):
        return parse_sets_excel(content)


class RebrickableClient:
    def __init__(self, api_key: str):
        if not api_key or not api_key.strip():
            raise ValueError("Clé API Rebrickable manquante.")
        self.api_key = api_key.strip()
        self._headers = {"Authorization": f"key {self.api_key}"}

    async def fetch_set(self, set_num: str) -> dict[str, Any]:
        async with httpx.AsyncClient(timeout=30.0) as client:
            set_resp = await client.get(
                f"{API_BASE}/lego/sets/{set_num}/",
                headers=self._headers,
            )
            if set_resp.status_code == 404:
                alt = (
                    set_num.rsplit("-", 1)[0] + "-1"
                    if "-" in set_num
                    else f"{set_num}-1"
                )
                if alt != set_num:
                    set_resp = await client.get(
                        f"{API_BASE}/lego/sets/{alt}/",
                        headers=self._headers,
                    )
                    set_num = alt
            set_resp.raise_for_status()
            data = set_resp.json()

            minifigs = 0
            mini_resp = await client.get(
                f"{API_BASE}/lego/sets/{set_num}/minifigs/",
                headers=self._headers,
                params={"page_size": 100},
            )
            if mini_resp.status_code == 200:
                payload = mini_resp.json()
                minifigs = sum(
                    int(item.get("quantity", 1) or 1)
                    for item in payload.get("results", [])
                )

        return {
            "set_num": data.get("set_num", set_num),
            "name": data.get("name"),
            "year": data.get("year"),
            "num_parts": data.get("num_parts") or 0,
            "num_minifigs": minifigs,
            "set_img_url": data.get("set_img_url"),
            "set_url": data.get("set_url"),
            "theme_id": data.get("theme_id"),
            "enriched": 1,
        }
